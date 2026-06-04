# -*- coding: utf-8 -*-
"""
HandVR 최종 발표용 PowerPoint 생성기.
출처: .omc/DESIGN_BRIEF.md, 보고서 5장(.omc/report_extracted.txt), Source/HandTrackingDemo/*, 에셋 추천 md.
16:9, 한글 폰트 Malgun Gothic. 산출물: docs/HandVR_최종발표.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FONT = "Malgun Gothic"

# 색상 팔레트 (차분한 실험실 톤)
C_BG      = RGBColor(0x10, 0x18, 0x28)   # 짙은 네이비
C_BG2     = RGBColor(0x17, 0x22, 0x38)
C_ACCENT  = RGBColor(0x4D, 0x9D, 0xE0)   # 블루
C_ACCENT2 = RGBColor(0x9B, 0x7E, 0xE0)   # 퍼플 (Threat)
C_GOLD    = RGBColor(0xF2, 0xC6, 0x4B)
C_WHITE   = RGBColor(0xF2, 0xF5, 0xFA)
C_GRAY    = RGBColor(0xB9, 0xC4, 0xD4)
C_DARKTX  = RGBColor(0x1A, 0x20, 0x2C)
C_GREEN   = RGBColor(0x4C, 0xC0, 0x7A)
C_RED     = RGBColor(0xE0, 0x5B, 0x5B)
C_PANEL   = RGBColor(0x1E, 0x2B, 0x44)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

SCR = r"C:/projects/VR/HandVR/.omc/screenshots"


def _set_font(run, size, color, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    # CJK/EastAsian 폰트도 명시
    try:
        rpr = run._r.get_or_add_rPr()
        from pptx.oxml.ns import qn
        ea = rpr.find(qn('a:ea'))
        if ea is None:
            ea = rpr.makeelement(qn('a:ea'), {})
            rpr.append(ea)
        ea.set('typeface', FONT)
        cs = rpr.find(qn('a:cs'))
        if cs is None:
            cs = rpr.makeelement(qn('a:cs'), {})
            rpr.append(cs)
        cs.set('typeface', FONT)
    except Exception:
        pass


def bg(slide, color=C_BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    return tb, tf


def add_para(tf, text, size, color, bold=False, italic=False, align=PP_ALIGN.LEFT,
             level=0, space_after=6, space_before=0, bullet=None, first=False, gap=None,
             before=None):
    if gap is not None:
        space_after = gap
    if before is not None:
        space_before = before
    p = tf.paragraphs[0] if first and tf.paragraphs[0].text == "" and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.level = level
    if space_after is not None:
        p.space_after = Pt(space_after)
    if space_before is not None:
        p.space_before = Pt(space_before)
    pre = ""
    if bullet == "dot":
        pre = "•  "
    elif bullet == "dash":
        pre = "–  "
    elif bullet == "num" and isinstance(level, int):
        pre = ""
    run = p.add_run()
    run.text = pre + text
    _set_font(run, size, color, bold=bold, italic=italic)
    return p


def header(slide, num, title, accent=C_ACCENT):
    # 좌측 액센트 바 + 제목 + 번호
    rect(slide, Inches(0), Inches(0), Inches(0.18), SH, accent)
    tb, tf = textbox(slide, Inches(0.55), Inches(0.32), Inches(11.2), Inches(0.95),
                     anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, title, 30, C_WHITE, bold=True, first=True, space_after=0)
    # 하단 구분선
    rect(slide, Inches(0.55), Inches(1.28), Inches(12.2), Pt(2), accent)
    # 슬라이드 번호
    tb2, tf2 = textbox(slide, Inches(12.3), Inches(0.32), Inches(0.85), Inches(0.6),
                       anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf2, f"{num:02d}", 16, C_GRAY, bold=True, align=PP_ALIGN.RIGHT, first=True)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def body_box(slide, lines, x=Inches(0.7), y=Inches(1.55), w=Inches(11.9), h=Inches(5.5),
             base_size=18):
    """lines: list of dict(text, size?, color?, bold?, level?, bullet?, gap?)"""
    tb, tf = textbox(slide, x, y, w, h)
    first = True
    for ln in lines:
        add_para(
            tf,
            ln["text"],
            ln.get("size", base_size),
            ln.get("color", C_GRAY),
            bold=ln.get("bold", False),
            italic=ln.get("italic", False),
            level=ln.get("level", 0),
            bullet=ln.get("bullet"),
            space_after=ln.get("gap", 7),
            space_before=ln.get("before", 0),
            first=first,
        )
        first = False
    return tb


def new_slide(color=C_BG):
    s = prs.slides.add_slide(BLANK)
    bg(s, color)
    return s


# ============================================================ Slide 1 — 표지
s = new_slide(C_BG)
# 장식 바
rect(s, Inches(0), Inches(2.55), SW, Inches(0.06), C_ACCENT)
rect(s, Inches(0), Inches(4.55), SW, Inches(0.02), C_PANEL)
tb, tf = textbox(s, Inches(1.0), Inches(1.0), Inches(11.3), Inches(1.4), anchor=MSO_ANCHOR.BOTTOM)
add_para(tf, "HandVR", 60, C_WHITE, bold=True, first=True, space_after=2)
tb, tf = textbox(s, Inches(1.0), Inches(2.7), Inches(11.3), Inches(1.6))
add_para(tf, "시각 우위성 기반 인지왜곡 VR 실험", 30, C_ACCENT, bold=True, first=True, space_after=6)
add_para(tf, "Rubber Hand Illusion · Visual Capture · Proprioceptive Drift", 17, C_GRAY,
         italic=True)
tb, tf = textbox(s, Inches(1.0), Inches(4.75), Inches(11.3), Inches(2.0))
add_para(tf, "가상현실 1팀", 24, C_GOLD, bold=True, first=True, space_after=8)
add_para(tf, "김희용 · 이종현 · 정윤아 · 이권수", 22, C_WHITE, space_after=14)
add_para(tf, "Meta Quest 3  ·  Unreal Engine 5.7  ·  C++  ·  OpenXR 핸드트래킹", 16, C_GRAY)
notes(s, "안녕하세요, 가상현실 1팀입니다. 저희 프로젝트 HandVR은 Quest 3 핸드트래킹과 실제 책상(패시브 햅틱)을 이용해, "
        "시각 우위성으로 인한 인지왜곡(고무손 착각·시각적 포착·고유수용감각 표류)을 1인 체험으로 검증하는 실험입니다. "
        "발표는 배경/가설 → 이론 3종 → 시스템·아키텍처 → 구현 현황 → 실험계획(8케이스) → 데이터·향후계획 순으로 진행합니다.")

# ============================================================ Slide 2 — 배경 & 목표
s = new_slide()
header(s, 2, "배경 & 목표")
body_box(s, [
    {"text": "배경", "size": 20, "color": C_ACCENT, "bold": True, "gap": 4},
    {"text": "VR에서 현실 소품(촉각)과 가상 오브젝트(시각)를 의도적으로 불일치시켜 감각 왜곡 과정을 실험적으로 탐구",
     "color": C_WHITE, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "핵심 측정: 시각 우위성이 촉각 인지에 미치는 영향", "color": C_WHITE, "bullet": "dot",
     "level": 1, "gap": 12},
    {"text": "세부 목표", "size": 20, "color": C_ACCENT, "bold": True, "gap": 4},
    {"text": "Quest 3 기반 1인 체험형 VR 실험 환경 구축 (실험실·책상·의자·조명)", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 4},
    {"text": "OpenXR 핸드트래킹으로 컨트롤러 없는 자연스러운 손 인터랙션", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 4},
    {"text": "패시브 햅틱스 + 물리 기반 렌더링으로 신체화(embodiment) 유도", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 4},
    {"text": "몰입감·실재감 변화 측정 및 감각왜곡 발생 여부 확인", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 4},
])
notes(s, "보고서 1장 목표 그대로입니다. 핵심은 보는 것과 만지는 것을 정해진 만큼만 어긋나게 만들어, 사람이 어느 감각을 더 믿는지를 "
        "정량적으로 보는 것입니다. 실제 책상을 같은 자리에 두는 패시브 햅틱이 신체화 유도의 전제 조건입니다.")

# ============================================================ Slide 3 — 핵심 가설
s = new_slide()
header(s, 3, "핵심 가설")
# 강조 박스
hb = rect(s, Inches(0.7), Inches(1.55), Inches(11.9), Inches(1.25), C_PANEL)
tb, tf = textbox(s, Inches(0.95), Inches(1.62), Inches(11.4), Inches(1.1), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf, "핵심 가설", 16, C_GOLD, bold=True, first=True, space_after=3)
add_para(tf, "시각 정보는 촉각 정보보다 우위에 있어, 충분히 정교한 VR 환경에서 사용자는 "
             "실제 자극과 불일치하는 감각왜곡을 경험한다.", 20, C_WHITE, bold=True)
body_box(s, [
    {"text": "세부 가설 (보고서 5장)", "size": 20, "color": C_ACCENT, "bold": True, "gap": 6},
    {"text": "공간 임계점 — 가상 손(시각)과 실제 손(고유수용감각)의 좌표 불일치가 커지면 "
             "신체소유감이 줄고, 임계값을 넘는 순간 인지왜곡이 중단", "color": C_WHITE,
     "bullet": "dot", "level": 1, "gap": 8},
    {"text": "촉각 왜곡 — 신체소유감이 활성화되면 실제 자극 부위와 시각 자극 부위가 달라도 "
             "뇌는 시각을 우선해 촉각왜곡이 발생", "color": C_WHITE, "bullet": "dot",
     "level": 1, "gap": 6},
], y=Inches(3.1), h=Inches(3.8))
notes(s, "가설은 두 갈래입니다. 첫째, 어긋남이 어느 선을 넘으면 착각이 깨진다(임계점). 둘째, 착각이 살아있는 상태에서는 시각이 "
        "촉각의 위치 판단까지 끌고 간다(촉각 왜곡). Phase 1이 첫째를, Phase 2가 둘째를 검증합니다.")

# ============================================================ Slide 4 — 이론1 RHI
s = new_slide()
header(s, 4, "이론 1 — 고무손 착각 (Rubber Hand Illusion)")
body_box(s, [
    {"text": "Botvinick & Cohen, 1998 — 후속 연구로 반복 검증", "size": 16, "color": C_GOLD,
     "italic": True, "gap": 10},
    {"text": "시각 자극과 촉각 자극이 동기화되면 가짜 손도 자신의 손처럼 느끼는 현상",
     "color": C_WHITE, "bullet": "dot", "gap": 12},
    {"text": "구현 (실험 1)", "size": 20, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "실제 추적 손을 숨기고, 책상 위 정적 가짜 손 위로 가상 붓이 0.5Hz(왕복 2초)로 쓰다듬음",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "조건: 자극 일치(보이는 위치=만지는 위치) / 불일치(다른 손가락 쪽 자극)",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "붓 왕복 1회마다 효과음 1회 (2초 클립, 이전 인스턴스 cut으로 누적 방지)",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
])
notes(s, "고무손 착각은 이 프로젝트의 출발점입니다. 보이는 손과 닿는 자극이 동기화되면 뇌가 가짜 손을 받아들입니다. "
        "코드에서는 실제 손을 숨기고, 책상 위 가짜 손을 붓이 일정 주기로 쓰다듬으며, 일치/불일치 조건을 프리셋으로 바꿉니다.")

# ============================================================ Slide 5 — 이론2 VC
s = new_slide()
header(s, 5, "이론 2 — 시각적 포착 (Visual Capture)")
body_box(s, [
    {"text": "여러 감각이 충돌할 때 사람은 일반적으로 시각 정보를 우선해 받아들이는 경향",
     "color": C_WHITE, "bullet": "dot", "gap": 5},
    {"text": "뇌가 신뢰도 높은 정보(시각)를 기준으로 다른 감각을 해석하는 과정으로 이해됨",
     "color": C_WHITE, "bullet": "dot", "gap": 12},
    {"text": "구현 (실험 2)", "size": 20, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "시각 손을 실제 손목 + 가변 공간오차(LateralOffsetCm, 예: 5 / 15 / 30 cm, 오른쪽)로 이동",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "책상 위 빨간 구(지름 ~8cm) 타겟을 \"보이는 손\"으로 터치",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "시각과 고유수용감각의 분리를 직접 체감", "color": C_GRAY, "bullet": "dot",
     "level": 1, "gap": 4},
])
notes(s, "시각적 포착은 손이 옆으로 점프해도 사용자는 보이는 손을 따라간다는 점을 보여줍니다. 빨간 구 타겟을 보이는 손으로 "
        "만지게 해 시각과 실제 손 위치의 분리를 직접 느끼게 합니다. 공간오차는 5~30cm로 가변입니다.")

# ============================================================ Slide 6 — 이론3 Drift
s = new_slide()
header(s, 6, "이론 3 — 고유수용감각 표류 (Proprioceptive Drift)")
body_box(s, [
    {"text": "시각의 영향으로 자신의 신체 위치를 실제와 다른 곳에 있다고 느끼는 현상 (RHI와 함께 관찰)",
     "color": C_WHITE, "bullet": "dot", "gap": 12},
    {"text": "구현 (실험 3)", "size": 20, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "Visual Capture + Y offset = base + 8·sin(2π·0.4·t) cm 의 느린 사인 흔들림",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "진폭 8cm, 주파수 0.4Hz (DriftAmplitudeCm / DriftHz)",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
    {"text": "실제 손목 위치에 파란 큐브 마커(4cm)를 띄워 \"진짜 손 위치\"를 시각화 — 시각 손과 분리",
     "color": C_GRAY, "bullet": "dot", "level": 1, "gap": 4},
])
notes(s, "표류 실험은 시각 포착을 시간축으로 확장합니다. 보이는 손을 천천히 흔들면 사용자의 위치 감각이 끌려옵니다. "
        "파란 마커는 실제 손목을 보여주어, 보이는 손과 진짜 손이 얼마나 벌어졌는지를 드러냅니다.")

# ============================================================ Slide 7 — 시스템 구성
s = new_slide()
header(s, 7, "시스템 구성")
# 표 형태: 2열 카드
rows = [
    ("엔진", "Unreal Engine 5.7.4 — VR 환경 · 렌더링 전체"),
    ("플러그인", "OculusXR Plugin 201.0 — Quest 3 핸드트래킹 / HMD 연동"),
    ("표준", "OpenXR (XR_EXT_hand_tracking) — 손 관절 추적 데이터 수신"),
    ("빌드", "Visual Studio 2026 / Android SDK 34 — PC 에디터 + Quest APK"),
    ("버전관리", "Git / GitHub — 소스코드 · 씬 에셋 공유"),
]
y = 1.7
for k, v in rows:
    rect(s, Inches(0.7), Inches(y), Inches(2.3), Inches(0.62), C_PANEL)
    tb, tf = textbox(s, Inches(0.75), Inches(y), Inches(2.2), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, k, 16, C_ACCENT, bold=True, first=True, align=PP_ALIGN.CENTER)
    tb, tf = textbox(s, Inches(3.15), Inches(y), Inches(9.4), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, v, 16, C_WHITE, first=True)
    y += 0.74
# 좌표/햅틱 강조
rect(s, Inches(0.7), Inches(y + 0.05), Inches(11.85), Inches(0.78), C_BG2)
tb, tf = textbox(s, Inches(0.9), Inches(y + 0.05), Inches(11.5), Inches(0.78), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf, "좌표 / 패시브 햅틱 — LocalFloor 원점, 책상 윗면 z=73, 책상앞 x=+30, 의자좌석 z=42 (가상=실제 정합)",
         16, C_GOLD, bold=True, first=True)
notes(s, "기술 스택은 보고서 2장 그대로입니다. 핵심은 OculusXR 플러그인으로 Quest 3의 손목 포즈를 직접 받고, OpenXR "
        "핸드트래킹으로 관절을 얹는 구조입니다. 좌표계는 LocalFloor 원점이라 사용자마다 책상 높이가 일정하게 잡혀 패시브 햅틱 "
        "정합이 유지됩니다.")

# ============================================================ Slide 8 — 아키텍처
s = new_slide()
header(s, 8, "아키텍처 — 클래스 구성 + FSM")
# 좌: 클래스, 우: FSM/손계층
tb, tf = textbox(s, Inches(0.7), Inches(1.55), Inches(6.6), Inches(5.4))
add_para(tf, "클래스 구성 (C++ 100%)", 18, C_ACCENT, bold=True, first=True, space_after=6)
classes = [
    ("AHandPawn", "VR 카메라 + 양손 OculusXR 핸드트래킹 + 손목 트리거 + 시각용 손 offset"),
    ("AExperimentManager", "런타임 FSM, 실험 props 스폰 / 조건 적용"),
    ("AHTDButton", "3D 손 터치 버튼 (박스 트리거 + 1초 쿨다운)"),
    ("ASurveyManager", "8문항 5점 리커트 + CSV 저장"),
    ("AWorldSetup", "절차적 방·책상·의자 (실측 사이즈 자동 fit)"),
]
for name, desc in classes:
    p = tf.add_paragraph(); p.space_after = Pt(7)
    r = p.add_run(); r.text = "▪ " + name + "  "; _set_font(r, 15, C_GOLD, bold=True)
    r2 = p.add_run(); r2.text = desc; _set_font(r2, 13.5, C_GRAY)
# 우측 패널
rect(s, Inches(7.55), Inches(1.55), Inches(5.05), Inches(5.4), C_PANEL)
tb, tf = textbox(s, Inches(7.8), Inches(1.7), Inches(4.6), Inches(5.1))
add_para(tf, "손 계층 (per hand)", 16, C_ACCENT, bold=True, first=True, space_after=4)
add_para(tf, "MotionController (실제 손목)", 13.5, C_WHITE, level=0, gap=2, bullet="dot")
add_para(tf, "→ WristTrigger (버튼 충돌)", 13, C_GRAY, level=1, gap=2)
add_para(tf, "→ HandOffset (시각 변위)", 13, C_GRAY, level=1, gap=2)
add_para(tf, "   → OculusXRHand (렌더 손)", 13, C_GRAY, level=1, gap=10)
add_para(tf, "상태 머신 (FSM)", 16, C_ACCENT, bold=True, space_after=4)
add_para(tf, "Intro → Start → MainMenu", 14, C_WHITE, gap=2, bullet="dot")
add_para(tf, "→ [Experiment | Survey | Results]", 13.5, C_GRAY, level=1, gap=2)
add_para(tf, "→ MainMenu", 13.5, C_GRAY, level=1, gap=8)
add_para(tf, "실험 종료 시 해당 조건 설문 자동 표시 → 메뉴 복귀", 13, C_GOLD, italic=True)
notes(s, "전부 C++로 작성했습니다. HandPawn은 실제 손목(MC) 아래에 변위용 HandOffset 노드를 두고, ExperimentManager가 "
        "이 노드만 밀어 시각 손을 옮깁니다. 실제 손은 그대로라 버튼 충돌은 진짜 손 기준입니다. 상태머신은 인트로부터 시작해 "
        "실험-설문-결과를 돌며, 실험이 끝나면 그 조건의 설문이 자동으로 뜹니다.")

# ============================================================ Slide 9 — 구현 현황
s = new_slide()
header(s, 9, "구현 현황")
# 완료 박스
rect(s, Inches(0.7), Inches(1.6), Inches(5.85), Inches(5.3), C_PANEL)
tb, tf = textbox(s, Inches(0.95), Inches(1.75), Inches(5.4), Inches(5.0))
add_para(tf, "✔ 완료", 20, C_GREEN, bold=True, first=True, space_after=8)
for t in [
    "핸드트래킹 파이프라인 (MC+OculusXRHand 계층, ConfidenceBehavior=None 끊김 방지)",
    "실험 4종 로직 (RHI / 시각포착 / 표류 / 망치위협)",
    "8문항 5점 리커트 설문 + CSV 자동 저장 (UTF-8 BOM)",
    "절차적 실험실 씬 (방·책상·의자·조명, 실측 사이즈 fit)",
    "붓 에셋 정합 (Blender→FBX→UE 임포트, 실측 비율, 효과음 임포트)",
]:
    add_para(tf, t, 14, C_WHITE, bullet="dot", gap=8)
# 진행 박스
rect(s, Inches(6.75), Inches(1.6), Inches(5.85), Inches(5.3), C_BG2)
tb, tf = textbox(s, Inches(7.0), Inches(1.75), Inches(5.4), Inches(5.0))
add_para(tf, "▶ 진행 중", 20, C_GOLD, bold=True, first=True, space_after=8)
for t in [
    "시간오차(Latency) 가변 적용 — 시각 손 N ms 지연 렌더(링버퍼)",
    "8개 실험 조건 자동 전환 로직 (조건별 자동 설문 + CSV)",
]:
    add_para(tf, t, 15, C_WHITE, bullet="dot", gap=10)
add_para(tf, "→ 두 기능 완성 시 보고서 5장 8케이스를 헤드셋을 벗지 않고 끝까지 자동 순회 가능",
         13.5, C_GRAY, italic=True, before=8)
notes(s, "완료된 것은 핸드트래킹, 실험 4종, 설문/CSV, 절차적 씬, 그리고 붓 에셋 정합입니다. 진행 중인 것은 두 가지인데, "
        "시각 손을 일정 시간 늦게 보여주는 시간오차 적용과, 8개 조건을 자동으로 순회시키는 전환 로직입니다. 이 둘이 붙으면 "
        "보고서 5장 8케이스를 헤드셋을 벗지 않고 끝까지 돌릴 수 있습니다.")

# ============================================================ Slide 10 — 8케이스 2페이즈 표
s = new_slide()
header(s, 10, "실험 계획 — 8케이스 2페이즈")
from pptx.enum.shapes import MSO_SHAPE
# Phase 1 라벨
tb, tf = textbox(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.4))
add_para(tf, "Phase 1 — 신체소유감 활성 임계값 도출", 16, C_ACCENT, bold=True, first=True)
t1 = s.shapes.add_table(5, 4, Inches(0.7), Inches(1.85), Inches(11.9), Inches(2.15)).table
p1 = [
    ["케이스", "공간오차", "시간오차", "실험 목적"],
    ["Case 1", "최소화", "최소화", "최상 동기화에서 신체소유감 측정"],
    ["Case 2", "가변 5~30cm", "최소화", "소유감이 사라지는 오차거리"],
    ["Case 3", "최소화", "가변 100~500ms", "소유감이 사라지는 오차시간"],
    ["Case 4", "가변 5~30cm", "가변 100~500ms", "시·공간 오차 공존 시 임계값"],
]
# Phase 2 라벨
tb, tf = textbox(s, Inches(0.7), Inches(4.15), Inches(12), Inches(0.4))
add_para(tf, "Phase 2 — 소유감 활성 상태의 촉각왜곡 검증", 16, C_ACCENT2, bold=True, first=True)
t2 = s.shapes.add_table(5, 4, Inches(0.7), Inches(4.55), Inches(11.9), Inches(2.15)).table
p2 = [
    ["케이스", "신체소유감", "촉각자극", "예상 결과"],
    ["Case 5", "활성", "불일치", "시각+표류가 촉각(현실)을 왜곡"],
    ["Case 6", "비활성", "불일치", "촉각왜곡 일어나지 않음"],
    ["Case 7", "활성", "시각자극만", "시각+표류가 촉각(현실)을 왜곡"],
    ["Case 8", "비활성", "시각자극만", "촉각왜곡 일어나지 않음"],
]


def fill_table(tbl, data, head_color):
    widths = [Inches(1.5), Inches(2.4), Inches(2.6), Inches(5.4)]
    for ci, w in enumerate(widths):
        tbl.columns[ci].width = w
    for ri, row in enumerate(data):
        tbl.rows[ri].height = Inches(0.42)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
            cell.margin_left = Pt(6); cell.margin_right = Pt(6)
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = head_color
            else:
                cell.fill.fore_color.rgb = C_PANEL if ri % 2 else C_BG2
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if ci < 3 else PP_ALIGN.LEFT
            r = p.add_run(); r.text = val
            _set_font(r, 12.5, C_WHITE if ri == 0 else (C_WHITE if ci == 0 else C_GRAY),
                      bold=(ri == 0 or ci == 0))


fill_table(t1, p1, C_ACCENT)
fill_table(t2, p2, C_ACCENT2)
# 각주
tb, tf = textbox(s, Inches(0.7), Inches(6.78), Inches(11.9), Inches(0.4))
add_para(tf, "* 4명이 각각 8케이스 수행. '일치'는 실제 자극부위 = 시각적 자극부위인 경우.",
         12, C_GRAY, italic=True, first=True)
notes(s, "보고서 5장의 핵심 표입니다. Phase 1은 어디까지 어긋나도 착각이 유지되는가라는 임계값을, Phase 2는 착각이 살아있을 때 "
        "시각이 촉각을 왜곡하는가를 봅니다. 4명이 각각 8케이스를 수행합니다. 일치는 실제 자극부위와 시각 자극부위가 동일한 경우를 뜻합니다.")

# ============================================================ Slide 11 — 조건 변인
s = new_slide()
header(s, 11, "조건 변인 (절충형 핵심)")
vars4 = [
    ("공간오차  LateralOffsetCm", "가변 프리셋 5 / 15 / 30 cm — 시각 손 +Y 이동", C_ACCENT),
    ("시간오차  LatencyMs  (신규)", "0 / 100 / 300 / 500 ms — 시각 손 지연 렌더(링버퍼) · '현실-가상 시간 정합' 담당", C_GOLD),
    ("자극 일치/불일치  StimulusMatched", "붓이 보이는 위치(일치) 또는 다른 손가락 쪽(불일치) 자극", C_ACCENT),
    ("신체소유감 유도  on / off", "붓 동기자극 유도 단계의 유무", C_ACCENT2),
]
y = 1.7
for title, desc, col in vars4:
    rect(s, Inches(0.7), Inches(y), Inches(0.14), Inches(1.0), col)
    tb, tf = textbox(s, Inches(1.0), Inches(y), Inches(11.5), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, title, 18, col, bold=True, first=True, space_after=2)
    add_para(tf, desc, 15, C_WHITE)
    y += 1.18
rect(s, Inches(0.7), Inches(y + 0.0), Inches(11.85), Inches(0.62), C_BG2)
tb, tf = textbox(s, Inches(0.9), Inches(y), Inches(11.5), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf, "절충형 = 현 4실험 위에 위 4변인을 '조건 프리셋'으로 얹어 8케이스를 재현 (재컴파일 불필요)",
         15, C_GOLD, bold=True, first=True)
notes(s, "절충형 설계의 요지입니다. 새로 실험을 만드는 대신, 이미 만든 4실험 위에 공간오차·시간오차·일치/불일치·소유감 on/off라는 "
        "변인 4개를 프리셋으로 얹어 8케이스를 그대로 재현합니다. 변인 값은 코드 EditAnywhere 프로퍼티와 조건 프리셋으로 관리해 "
        "재컴파일 없이 바꿀 수 있습니다.")

# ============================================================ Slide 12 — 자극 & 인터랙션
s = new_slide()
header(s, 12, "자극 & 인터랙션")
exps = [
    ("실험 1 · RHI", "가짜 손 위 가상 붓 0.5Hz 왕복(왕복 2초), 일치/불일치, 왕복당 효과음 1회", C_ACCENT),
    ("실험 2 · 시각포착", "시각 손 +공간오차, 책상 위 빨간 구(~8cm) 타겟을 보이는 손으로 터치", C_ACCENT),
    ("실험 3 · 표류", "VC + 8cm·0.4Hz 사인 흔들림, 실제 손목에 파란 큐브 마커(4cm)", C_ACCENT),
    ("실험 4 · 망치위협", "lift → strike → impact → return 4단계 반복, 충격 순간 효과음 1회", C_ACCENT2),
]
y = 1.65
for title, desc, col in exps:
    rect(s, Inches(0.7), Inches(y), Inches(2.55), Inches(0.78), col)
    tb, tf = textbox(s, Inches(0.75), Inches(y), Inches(2.45), Inches(0.78), anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, title, 15, C_DARKTX, bold=True, first=True, align=PP_ALIGN.CENTER)
    tb, tf = textbox(s, Inches(3.4), Inches(y), Inches(9.2), Inches(0.78), anchor=MSO_ANCHOR.MIDDLE)
    add_para(tf, desc, 14.5, C_WHITE, first=True)
    y += 0.92
tb, tf = textbox(s, Inches(0.7), Inches(y + 0.05), Inches(11.9), Inches(1.4))
add_para(tf, "공통", 16, C_GOLD, bold=True, first=True, space_after=4)
add_para(tf, "모든 시각 효과는 즉시(스냅) 적용/해제, 프롭 충돌 없음(매 틱 위치 갱신)", 14, C_GRAY,
         bullet="dot", gap=4)
add_para(tf, "사운드: 프롭 부착 UAudioComponent, 이벤트마다 Play()로 이전 인스턴스 cut(누적 방지), 2초 클립",
         14, C_GRAY, bullet="dot", gap=4)
notes(s, "자극은 4명 모두에게 동일하게 들어가도록 고정했습니다. 붓은 정해진 손가락 위에서 일정 주기로 왕복하고, 망치는 "
        "들어올림-내려치기-충격-복귀 4단계를 반복합니다. 충격이나 왕복 순간에만 2초 효과음이 한 번 울리도록 엣지 검출로 누적을 "
        "막았습니다. 효과는 페이드 없이 즉시 적용됩니다.")

# ============================================================ Slide 13 — 설문 8문항
s = new_slide()
header(s, 13, "설문 — 8문항 5점 리커트")
tb, tf = textbox(s, Inches(0.7), Inches(1.4), Inches(11.9), Inches(0.4))
add_para(tf, "각 케이스 종료 시 자동 표시 · 1점(전혀 아니다) ~ 5점(매우 그렇다)", 14, C_GOLD,
         italic=True, first=True)
t = s.shapes.add_table(9, 3, Inches(0.7), Inches(1.85), Inches(11.9), Inches(5.0)).table
qs = [
    ["#", "측정항목 · 문항", "분석 용도"],
    ["1", "공간일치감 — 가상 손 위치가 실제 내 손 위치와 비슷했는가?", "공간 허용범위"],
    ["2", "시간일치감 — 본 자극과 실제 촉각이 동시에 일어난 것 같았는가?", "시간 동기화 허용범위"],
    ["3", "촉각왜곡 — 실제 자극부위가 아니라 본 위치에서 촉각이 느껴졌는가?", "왜곡 발생 여부"],
    ["4", "자극위치판단 — 자극 위치를 시각적으로 본 위치 기준으로 판단했는가?", "시각의 위치판단 영향"],
    ["5", "시각우위성 — 몸의 느낌보다 보이는 정보를 더 신뢰했는가?", "시각 우위성"],
    ["6", "위화감 — 가상 손과 실제 손 사이 어색함/위화감을 느꼈는가?", "몰입도"],
    ["7", "현실감 — 자극 상황이 실제 내 손에 일어난 일처럼 느껴졌는가?", "실재감"],
    ["8", "위협감 — 가상 자극이 실제로 내 손을 위협한다고 느껴졌는가?", "위협 설득력"],
]
widths = [Inches(0.6), Inches(8.0), Inches(3.3)]
for ci, w in enumerate(widths):
    t.columns[ci].width = w
for ri, row in enumerate(qs):
    t.rows[ri].height = Inches(0.55)
    for ci, val in enumerate(row):
        cell = t.cell(ri, ci)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
        cell.margin_left = Pt(7); cell.margin_right = Pt(6)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_GOLD if ri == 0 else (C_PANEL if ri % 2 else C_BG2)
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val
        _set_font(r, 12.5, C_DARKTX if ri == 0 else (C_WHITE if ci <= 1 else C_GRAY),
                  bold=(ri == 0 or ci == 0))
notes(s, "설문은 보고서 5장의 8문항을 그대로 옮겼습니다. 공간/시간 일치감, 촉각왜곡, 자극위치판단, 시각우위성, 위화감, 현실감, "
        "위협감 순입니다. 케이스별 평균 점수를 비교해 공간·시간·촉각 차이가 소유감과 촉각왜곡에 미치는 영향을 분석합니다. "
        "(구현 코드는 신체소유감 문항을 맨 앞에 추가해 운용상 9문항으로 운영하기도 합니다.)")

# ============================================================ Slide 14 — 데이터 로깅
s = new_slide()
header(s, 14, "데이터 로깅 & 분석")
body_box(s, [
    {"text": "조건별 CSV 자동 저장", "size": 19, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "Saved/SurveyResults_<timestamp>.csv  (UTF-8 BOM, Excel 호환)", "color": C_WHITE,
     "bullet": "dot", "level": 1, "gap": 4},
    {"text": "한 행 = 조건 메타(실험종류 · LateralOffsetCm · LatencyMs · Matched) + 8문항 점수 + 타임스탬프",
     "color": C_WHITE, "bullet": "dot", "level": 1, "gap": 12},
    {"text": "사고 방지 (보고서 4.6)", "size": 19, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "설문 미완료 시 중간값(3) 자동 기록 후 진행", "color": C_GRAY, "bullet": "dot",
     "level": 1, "gap": 4},
    {"text": "이벤트 즉시 flush — 앱 중단에도 직전까지 데이터 보존", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 12},
    {"text": "분석", "size": 19, "color": C_ACCENT, "bold": True, "gap": 5},
    {"text": "케이스별 평균 점수 비교로 공간·시간·촉각 변인의 효과를 정량화", "color": C_GRAY,
     "bullet": "dot", "level": 1, "gap": 4},
])
notes(s, "데이터는 조건이 끝날 때마다 한 행씩 CSV로 즉시 기록합니다. 설문이 안 뜨는 사고가 나도 중간값을 자동 기록하고 넘어가며, "
        "모든 이벤트를 발생 즉시 파일에 써서 중간에 멈춰도 직전까지는 보존됩니다. 분석은 케이스별 평균 점수 비교로 진행합니다.")

# ============================================================ Slide 15 — 인트로 스토리 & 흐름
s = new_slide()
header(s, 15, "인트로 스토리 & 사용자 흐름")
body_box(s, [
    {"text": "인트로 내러티브 (앱 시작, 시작버튼 전) — 신체소유감 유도 사전조건 형성", "size": 18,
     "color": C_ACCENT, "bold": True, "gap": 6},
    {"text": "\"지각 실험실에 오신 것을 환영합니다. 의자에 앉아 손을 책상 위에 올려 주세요.\"",
     "color": C_WHITE, "bullet": "dot", "level": 1, "gap": 3},
    {"text": "\"곧 보이는 것과 느껴지는 것이 어긋날 수 있습니다. 판단하지 말고 관찰하세요.\"",
     "color": C_WHITE, "bullet": "dot", "level": 1, "gap": 3},
    {"text": "\"준비되면 [시작]을 눌러 주세요.\"", "color": C_WHITE, "bullet": "dot",
     "level": 1, "gap": 14},
], y=Inches(1.55), h=Inches(2.7))
# 흐름 다이어그램
flow = ["Intro", "시작", "메인메뉴", "실험\n(조건 적용)", "조건 설문\n자동 표시", "메뉴 복귀"]
x = 0.7
yy = 4.55
bw = 1.78
for i, step in enumerate(flow):
    col = C_ACCENT if i not in (3, 4) else C_GOLD
    rect(s, Inches(x), Inches(yy), Inches(bw), Inches(0.95), C_PANEL, line=col)
    tb, tf = textbox(s, Inches(x), Inches(yy), Inches(bw), Inches(0.95), anchor=MSO_ANCHOR.MIDDLE)
    for j, line in enumerate(step.split("\n")):
        add_para(tf, line, 14, C_WHITE if i not in (3, 4) else C_GOLD, bold=True,
                 align=PP_ALIGN.CENTER, first=(j == 0), gap=0)
    if i < len(flow) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.02), Inches(yy + 0.32),
                                Inches(0.28), Inches(0.3))
        ar.fill.solid(); ar.fill.fore_color.rgb = C_ACCENT; ar.line.fill.background()
        ar.shadow.inherit = False
    x += bw + 0.32
tb, tf = textbox(s, Inches(0.7), Inches(5.8), Inches(11.9), Inches(1.3))
add_para(tf, "메뉴: 실험 1~4 / 설문 / 결과 / 종료  +  실험 중 '중단' 버튼 → 즉시 메뉴 복귀 (반복 / 종료)",
         15, C_GRAY, first=True)
notes(s, "사용자 경험은 인트로 내러티브로 시작합니다. 차분한 실험실 안내로 관찰자 모드를 만들어 신체소유감 유도의 전제를 깝니다. "
        "이후 메뉴에서 실험을 고르면 조건이 즉시 적용되고, 끝나면 그 조건의 설문이 자동으로 떠 응답 후 메뉴로 돌아옵니다.")

# ============================================================ Slide 16 — 역할분담 & 일정
s = new_slide()
header(s, 16, "역할 분담 & 일정 (보고서 6장)")
# 완료
tb, tf = textbox(s, Inches(0.7), Inches(1.5), Inches(5.9), Inches(0.4))
add_para(tf, "완료 사항", 18, C_GREEN, bold=True, first=True)
done = [
    ("김희용", "제안서·보고서, 실험 절차 설계 및 설문 문항 구성"),
    ("이종현", "GitHub 협업환경, 핸드트래킹 파이프라인 구축"),
    ("정윤아", "발표자료 제작, 가상현실 씬 구성"),
    ("이권수", "아이디어 제시, 제안서·보고서"),
]
tb, tf = textbox(s, Inches(0.7), Inches(1.95), Inches(5.9), Inches(4.9))
first = True
for name, role in done:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(10)
    r = p.add_run(); r.text = name + "  "; _set_font(r, 15, C_GOLD, bold=True)
    r2 = p.add_run(); r2.text = role; _set_font(r2, 13.5, C_GRAY)
# 계획
tb, tf = textbox(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(0.4))
add_para(tf, "계획 사항", 18, C_GOLD, bold=True, first=True)
plan = [
    ("김희용", "설문 UI·데이터 기록, 실험 진행 및 결과 발표"),
    ("이종현", "현실-가상 좌표/시간 정합, 8조건 자동 전환 로직"),
    ("정윤아", "실험 도구 모델링·텍스처·조명·최적화, 발표자료"),
    ("이권수", "사운드 시스템, 데이터 분석 및 최종 보고서"),
]
tb, tf = textbox(s, Inches(6.85), Inches(1.95), Inches(5.9), Inches(4.9))
first = True
for name, role in plan:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(10)
    r = p.add_run(); r.text = name + "  "; _set_font(r, 15, C_ACCENT, bold=True)
    r2 = p.add_run(); r2.text = role; _set_font(r2, 13.5, C_GRAY)
# 구분선
rect(s, Inches(6.72), Inches(1.5), Pt(1.5), Inches(4.5), C_PANEL)
notes(s, "역할은 보고서 6장 그대로입니다. 완료 단계에서 설계·핸드트래킹·씬·아이디어를 각각 맡았고, 계획 단계에서는 "
        "설문/데이터(김희용), 좌표·시간 정합 및 조건 자동전환(이종현), 모델링·최적화(정윤아), 사운드·분석(이권수)으로 나눕니다.")

# ============================================================ Slide 17 — 데모 스크린샷
s = new_slide()
header(s, 17, "데모 — 실험실 씬 스크린샷")
shots = [
    (os.path.join(SCR, "02_three_quarter.png"), "씬 3/4 뷰 — 책상 · 의자 · 룸"),
    (os.path.join(SCR, "04_desk_closeup.png"), "책상 클로즈업 — 프롭 배치"),
    (os.path.join(SCR, "01_seated_pov.png"), "착석 시점 — z=73 책상 정합"),
]
img_w = Inches(3.9)
img_h = Inches(2.45)
gap = Inches(0.18)
total = img_w * 3 + gap * 2
x0 = Emu(int((SW - total) / 2))
y0 = Inches(2.0)
placed = 0
for path, cap in shots:
    cx = Emu(int(x0) + placed * (int(img_w) + int(gap)))
    if os.path.exists(path):
        try:
            s.shapes.add_picture(path, cx, y0, width=img_w, height=img_h)
        except Exception:
            rect(s, cx, y0, img_w, img_h, C_PANEL, line=C_ACCENT)
    else:
        ph = rect(s, cx, y0, img_w, img_h, C_PANEL, line=C_ACCENT)
        tbp, tfp = textbox(s, cx, y0, img_w, img_h, anchor=MSO_ANCHOR.MIDDLE)
        add_para(tfp, "[스크린샷 자리]", 16, C_GRAY, align=PP_ALIGN.CENTER, first=True)
    # 캡션
    tbc, tfc = textbox(s, cx, Emu(int(y0) + int(img_h) + Emu(50000)), img_w, Inches(0.5))
    add_para(tfc, cap, 13, C_WHITE, align=PP_ALIGN.CENTER, first=True)
    placed += 1
tb, tf = textbox(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(1.4))
add_para(tf, "절차적 실험실 씬 — 책상(z=73) · 의자(좌석 z=42) · 가짜 손/프롭 실측 배치", 15, C_GRAY,
         bullet="dot", first=True, gap=5)
add_para(tf, "헤드리스 캡처 환경(이 PC의 OpenXR/TDR 제약)으로 화면은 어둡게 보임 — 시각 최종검증은 Quest 3 실기기(월요일 팀 실험)",
         14, C_GOLD, italic=True, bullet="dot")
notes(s, "현재 빌드의 실험실 씬입니다. 실측 사이즈로 책상과 의자가 배치되고, 책상 위에 가짜 손과 프롭이 올라갑니다. 헤드리스 캡처 "
        "환경(이 PC의 OpenXR/TDR 제약) 한계로 화면은 다소 어둡지만, 시각 최종검증은 Quest 3 실기기에서 진행합니다.")

# ============================================================ Slide 18 — 결론 & 향후계획
s = new_slide()
header(s, 18, "결론 & 향후 계획")
# 결론
tb, tf = textbox(s, Inches(0.7), Inches(1.55), Inches(11.9), Inches(2.6))
add_para(tf, "결론", 19, C_GREEN, bold=True, first=True, space_after=5)
for t in [
    "시각 우위성 인지왜곡(RHI·VC·Drift)을 Quest 3 + 핸드트래킹 + 패시브 햅틱으로 1인 검증하는 파이프라인 구축",
    "4실험 + 절충형 조건 변인으로 보고서 5장 8케이스를 재현 가능한 구조 확보",
    "C++ FSM · 설문 · CSV · 절차적 씬 · 붓 에셋 정합까지 완료",
]:
    add_para(tf, t, 14.5, C_WHITE, bullet="dot", gap=5)
# 향후
tb, tf = textbox(s, Inches(0.7), Inches(4.25), Inches(11.9), Inches(2.4))
add_para(tf, "향후 계획", 19, C_GOLD, bold=True, first=True, space_after=5)
for t in [
    "시간오차(Latency) 가변 적용 + 8조건 자동 전환 완성",
    "실험 도구 모델링/텍스처·조명 최적화, 사운드 시스템 마무리",
    "Quest 3 실기기 실험(4인 × 8케이스) → CSV 수집 → 케이스별 분석 → 최종 보고서",
]:
    add_para(tf, t, 14.5, C_GRAY, bullet="dot", gap=5)
# 한계
rect(s, Inches(0.7), Inches(6.55), Inches(11.85), Inches(0.62), C_BG2)
tb, tf = textbox(s, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
add_para(tf, "한계(정직하게): 헤드리스 SceneCapture / Meta XR Simulator 불안정 → 시각 최종검증은 Quest 3 실기기",
         13.5, C_GOLD, italic=True, first=True)
notes(s, "정리하면, 인지왜곡을 1인 체험으로 검증하는 파이프라인을 C++로 구축했고, 절충형 변인 설계로 8케이스를 재현할 구조를 "
        "갖췄습니다. 남은 일은 시간오차와 조건 자동전환을 붙이고, 모델링·사운드를 마무리한 뒤, 실기기에서 4명×8케이스 데이터를 모아 "
        "분석하는 것입니다. 시각 검증은 이 PC의 시뮬레이터 한계로 실기기에서 마무리합니다. 감사합니다.")

# ============================================================ 저장
out = r"C:/projects/VR/HandVR/docs/HandVR_최종발표.pptx"
prs.save(out)
print("SAVED", out, "slides=", len(prs.slides._sldIdLst))
